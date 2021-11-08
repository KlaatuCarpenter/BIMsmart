from openpyxl import load_workbook
import random
import pandas as pd

scheduleOfValuesPath = "../input_data_used_for_valuation/schedule_of_values.xlsx"
paymentProgressPath = "../input_data_used_for_valuation"
paymentID = 3

previous = 'Previous settlement period'
current = 'Current settlement period'
total = 'Total progress'

    
if __name__ == "__main__":
    if paymentID == 0:
        # @dev this should be proceeded if its the first payment
        # FIRST PAYMENT
        # This code adds additional columns nedded to proceed further transactions

        scheduleOfValues = load_workbook(scheduleOfValuesPath, data_only=True)['Summary']

        if len(scheduleOfValues.tables.items()) > 1:
            print('Error during importing the schedule of value table. More than one table in worksheet.')
        
        data_boundary = scheduleOfValues.tables.items()[0][1]
        data = scheduleOfValues[data_boundary]
        # extract the data 
        # the inner list comprehension gets the values for each cell in the table
        content = [[cell.value for cell in ent] 
                for ent in data
            ]
            
        header = content[0]
        indexes = []
        for item in content:
            indexes.append(item[0])
            # to avoid redundancy in index and content:
            item.pop(0)
        # remove index which is in header
        indexes.pop(0)
        
        # the contents ... excluding the header            
        rest = content[1:]
        
        # create dataframe with the column names
        # and pair table name with dataframe
        schedule = pd.DataFrame(rest, columns = pd.MultiIndex.from_product( [['Contract'], header]), index = indexes)

        # crate table to measure progress in settlement period
        initialValues = []
        for i in range(len(indexes)):
            initialValues.append([0,0,0,0,0,0])
        
        aboveHeader = [[previous, current, total], ['Count', 'Value'] ]

        progress = pd.DataFrame(initialValues, columns = pd.MultiIndex.from_product(aboveHeader), index = indexes)

        paymentTable = pd.concat([schedule, progress], axis=1)

        # write to file 
        paymentTable.to_excel(f'{paymentProgressPath}/payment{paymentID}_progress.xlsx')

    else:
        # Proceess payment evaluation
        paymentProgress = pd.read_excel(f'{paymentProgressPath}/payment{paymentID-1}_progress.xlsx', header = [0, 1], index_col = 0)
        
        for item in paymentProgress.index:
            # put previous total progress from last payment to previous settlement period
            paymentProgress.loc[item,(previous, 'count')] = paymentProgress.loc[item,(total, 'count')]
            paymentProgress.loc[item,(previous, 'value')] = paymentProgress.loc[item,(total, 'value')]

            # for the example a complicated computations with AI and on site captured data is just random ...
            toDo = paymentProgress.loc[item,('Contract', 'Count')] - paymentProgress.loc[item,(previous, 'count')]
            paymentProgress.loc[item,(current, 'count')] = random.randrange(toDo)
            paymentProgress.loc[item,(current, 'value')] = paymentProgress.loc[item,(current, 'count')] * paymentProgress.loc[item,('Contract', 'Unit price')]

            # calculate total progress
            paymentProgress.loc[item,(total, 'count')] = paymentProgress.loc[item,(current, 'count')] + paymentProgress.loc[item,(previous, 'count')]
            paymentProgress.loc[item,(total, 'value')] = paymentProgress.loc[item,(current, 'value')] + paymentProgress.loc[item,(previous, 'value')]

        # calculate total progress value
        currentProgressValue = paymentProgress[(current, 'value')].sum()
        totalProgressValue = paymentProgress[(total, 'value')].sum()
        
        # write to file 
        paymentProgress.to_excel(f'{paymentProgressPath}/payment{paymentID}_progress.xlsx')









    
     

