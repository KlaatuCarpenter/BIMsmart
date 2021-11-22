from openpyxl import load_workbook
import pandas as pd
import requests
import common.util
import os
from resources.solution_used_for_progress_evaluation import progress_algorythm
from resources.create_metadata import Metadata


# base url for Fleek's IPFS Gateway
base_url = 'https://ipfs.fleek.co/ipfs/'
# dict to convenient name cells in payment progress
dict = {
    'previous': 'Previous settlement period',
    'current': 'Current settlement period',
    'total': 'Total progress',
    'count_str': 'Count',
    'value_str': 'Value'
}


class Update:

    def __init__(self, input):
        self.request_data = input.get('data')
        self.id = self.request_data['paymentID']
        if self.validate_request_data():
            self.evaluate_progress_payment()
        else:
            self.result_error('No data provided')

    def validate_request_data(self):
        if self.request_data is None:
            return False
        if self.request_data == {}:
            return False
        return True

    def evaluate_progress_payment(self):
        if self.id == 1:
            payment_progress_0 = self.create_initial_payment_progress()
            inputFile = pd.read_excel(
                payment_progress_0, header=[0, 1], index_col=0)
            os.remove(payment_progress_0)
        else:
            req = requests.get(
                f"{base_url}{self.request_data['CID_previousPaymentProgress']}")
            with open('scheduleprogress.xlsx', 'wb') as f:
                f.write(req.content)
            # Proceess payment evaluation
            inputFile = pd.read_excel(
                'scheduleprogress.xlsx', header=[0, 1], index_col=0)
            os.remove('scheduleprogress.xlsx')

        for item in inputFile.index:
            # put previous total progress from last payment to previous settlement period
            inputFile.loc[item, (dict['previous'], dict['count_str'])
                          ] = inputFile.loc[item, (dict['total'], dict['count_str'])]
            inputFile.loc[item, (dict['previous'], dict['value_str'])
                          ] = inputFile.loc[item, (dict['total'], dict['value_str'])]

            # for the example a complicated computations with AI and on site captured data is processed
            # with the function named "progress_algorythm" ...
            toDo = inputFile.loc[item, ('Contract', dict['count_str'])] - \
                inputFile.loc[item, (dict['previous'], dict['count_str'])]
            inputFile.loc[item, (dict['current'], dict['count_str'])
                          ] = progress_algorythm(toDo)
            inputFile.loc[item, (dict['current'], dict['value_str'])] = inputFile.loc[item, (
                dict['current'], dict['count_str'])] * inputFile.loc[item, ('Contract', 'Unit price')]

            # calculate total progress
            inputFile.loc[item, (dict['total'], dict['count_str'])] = inputFile.loc[item, (
                dict['current'], dict['count_str'])] + inputFile.loc[item, (dict['previous'], dict['count_str'])]
            inputFile.loc[item, (dict['total'], dict['value_str'])] = inputFile.loc[item, (
                dict['current'], dict['value_str'])] + inputFile.loc[item, (dict['previous'], dict['value_str'])]

        # calculate total progress value
        self.currentProgressValue = int(
            inputFile[(dict['current'], dict['value_str'])].sum())
        self.totalProgressValue = int(
            inputFile[(dict['total'], dict['value_str'])].sum())

        # write to file
        payment_progress_filename = f"{self.id}_payment_progress.xlsx"
        inputFile.to_excel(payment_progress_filename)

        try:
            res = common.util.save_file_in_ipfs(
                payment_progress_filename, payment_progress_filename)
            self.CID_payment_progress = res['ResponseMetadata']['HTTPHeaders']['x-fleek-ipfs-hash-v0']
            # check CID of solution used for progress evaluation
            res_solution = common.util.save_file_in_ipfs(
                f'{self.id}_solution_used_for_progress_evaluation.py', 'resources/solution_used_for_progress_evaluation.py')
            self.CID_solution = res_solution['ResponseMetadata']['HTTPHeaders']['x-fleek-ipfs-hash-v0']
            os.remove(payment_progress_filename)
            self.result_success()
        except Exception as e:
            self.result_error(e)
            os.remove(payment_progress_filename)

    # if it is a first payment it creates the progress payment file with initial values equals to 0 to work on
    def create_initial_payment_progress(self):
        req = requests.get(
            f"{base_url}{self.request_data['CID_scheduleOfValues']}")
        with open('schedule.xlsx', 'wb') as f:
            f.write(req.content)
        scheduleOfValues = load_workbook(
            'schedule.xlsx', data_only=True)['Summary']
        os.remove('schedule.xlsx')

        if len(scheduleOfValues.tables.items()) > 1:
            self.result_error(
                'Error during importing the schedule of value table. More than one table in worksheet.')

        data_boundary = scheduleOfValues.tables.items()[0][1]
        data = scheduleOfValues[data_boundary]
        # extract the data
        # the inner list comprehension gets the values for each cell in the table
        content = [[cell.value for cell in ent]
                   for ent in data
                   ]

        header = content[0]
        indexes = []
        for item in content:
            indexes.append(item[0])
            # to avoid redundancy in index and content:
            item.pop(0)
        # remove index which is in header
        indexes.pop(0)

        # the contents ... excluding the header
        rest = content[1:]

        # create dataframe with the column names
        # and pair table name with dataframe
        schedule = pd.DataFrame(rest, columns=pd.MultiIndex.from_product(
            [['Contract'], header]), index=indexes)

        # crate table to measure progress in settlement period
        initialValues = []
        for i in range(len(indexes)):
            initialValues.append([0, 0, 0, 0, 0, 0])
        aboveHeader = [[dict['previous'], dict['current'], dict['total']], [
            dict['count_str'], dict['value_str']]]
        progress = pd.DataFrame(
            initialValues, columns=pd.MultiIndex.from_product(aboveHeader), index=indexes)
        paymentTable = pd.concat([schedule, progress], axis=1)

        # write to file
        payment_progress_filename_0 = "0_payment_progress.xlsx"
        paymentTable.to_excel(payment_progress_filename_0)
        return payment_progress_filename_0

    def result_success(self):
        self.result = {
            'name': self.request_data['name'],
            'paymentID': self.request_data['paymentID'],
            'CID_listOfElementsAndGUIDs': self.request_data['CID_listOfElementsAndGUIDs'],
            'CID_asBuiltBIM': self.request_data['CID_asBuiltBIM'],
            'CID_scheduleOfValues': self.request_data['CID_scheduleOfValues'],
            'CID_rawProgressData': self.request_data['CID_rawProgressData'],
            'CID_solutionUsedForProgressEvaluation': self.CID_solution,
            'value': self.currentProgressValue,
            'total_contract_progress': self.totalProgressValue,
            'CID_currentPaymentProgress': self.CID_payment_progress,
            'CID_previousPaymentProgress': self.request_data['CID_previousPaymentProgress'],
            'statusCode': 200,
        }
        NFT_metadata = Metadata(self.result)
        try:
            self.result['NFT_URI'] = NFT_metadata.CID
        except AttributeError:
            self.result_error(NFT_metadata.error)

    def result_error(self, error):
        self.result = {
            'paymentID': self.id,
            'status': 'errored',
            'error': f'There was an error: {error}',
            'statusCode': 500,
        }
